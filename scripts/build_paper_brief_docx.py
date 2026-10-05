"""Build the condensed paper plan as a DOCX artifact."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "paper" / "docs" / "zhasyrynsoz_condensed_paper_plan.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text: str, bold: bool = False) -> None:
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9.5)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_border(cell)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.width = Inches(widths[i])
        set_cell_shading(cell, "1F2937")
        set_cell_text(cell, header, bold=True)
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)

    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].width = Inches(widths[i])
            set_cell_text(cells[i], value)

    doc.add_paragraph()


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ["Title", "Heading 1", "Heading 2"]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(0, 0, 0)
    styles["Title"].font.size = Pt(20)
    styles["Heading 1"].font.size = Pt(14)
    styles["Heading 2"].font.size = Pt(12)


def build() -> None:
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("Жасырын Сөз Research Paper Plan")

    subtitle = doc.add_paragraph()
    subtitle.add_run("Condensed academic roadmap for a Kazakh lexical semantics paper").italic = True

    doc.add_heading("Purpose", level=1)
    doc.add_paragraph(
        "This document compresses the broader project plan into a paper-first research plan. "
        "The game is treated as infrastructure for evaluating Kazakh word-level semantic "
        "representations, not as the main deliverable."
    )

    doc.add_heading("Central Claim", level=1)
    doc.add_paragraph(
        "Kazakh has limited public resources for evaluating lexical semantic similarity. "
        "Жасырын Сөз can provide a reproducible vocabulary, a human annotation task, and "
        "an interpretable testbed for comparing embedding models against native-speaker judgments."
    )

    doc.add_heading("Research Question", level=1)
    doc.add_paragraph(
        "How well do available static and transformer-based embedding models capture Kazakh "
        "word-level semantic similarity, and where do their rankings diverge from human intuition?"
    )

    doc.add_heading("Expected Contributions", level=1)
    add_bullets(
        doc,
        [
            "A Kazakh semantic word-guessing game suitable for controlled model evaluation.",
            "A documented curated core vocabulary separated from noisy candidate vocabulary.",
            "A native-speaker word-pair similarity annotation set.",
            "A comparison of FastText and multilingual transformer embeddings for Kazakh word similarity.",
            "An error analysis covering morphology, loanwords, topical association, and polysemy.",
        ],
    )

    doc.add_heading("Minimum Publishable Study", level=1)
    add_table(
        doc,
        ["Component", "Minimum version", "Why it matters"],
        [
            ["Vocabulary", "Curated core words from the game list", "Defines the evaluation universe and avoids noisy artifacts."],
            ["Human data", "150-300 native-speaker scored word pairs", "Provides the reference signal for model comparison."],
            ["Models", "FastText wiki.kk, FastText cc.kk, multilingual-e5, bge-m3", "Covers static and transformer-based baselines."],
            ["Metrics", "Spearman rho, Pearson r, neighbor-quality review", "Measures both rank agreement and qualitative usefulness."],
            ["Ethics", "Consent and anonymization for any gameplay logs", "Keeps optional behavioral data publishable."],
        ],
        [1.35, 2.2, 3.1],
    )

    doc.add_heading("Method Pipeline", level=1)
    add_table(
        doc,
        ["Step", "Output"],
        [
            ["Export vocabulary", "paper/data/vocabulary_export.csv"],
            ["Sample annotation pairs", "paper/data/annotation_pairs_seed.csv"],
            ["Human annotation", "Scores from 0 to 10 for semantic similarity"],
            ["Precompute embeddings", "One vector table per model for the same vocabulary"],
            ["Evaluate models", "Correlation table and top-neighbor error categories"],
            ["Write results", "Results, analysis, limitations, and release notes"],
        ],
        [2.05, 4.6],
    )

    doc.add_heading("Paper Structure", level=1)
    add_bullets(
        doc,
        [
            "Introduction: motivation, research gap, and contributions.",
            "Related Work: semantic guessing games, lexical similarity datasets, and Kazakh NLP.",
            "System: game architecture and similarity feedback mechanism.",
            "Dataset: vocabulary construction, annotation protocol, and quality control.",
            "Models and Evaluation: embedding models, preprocessing, and metrics.",
            "Results: correlations and neighbor-quality comparison.",
            "Error Analysis: model failures and linguistic patterns.",
            "Limitations and Ethics: dataset size, annotator pool, consent, and anonymization.",
        ],
    )

    doc.add_heading("Immediate Next Actions", level=1)
    add_table(
        doc,
        ["Priority", "Action", "Status"],
        [
            ["1", "Review the 150 seed annotation pairs and remove confusing or culturally narrow pairs.", "Ready"],
            ["2", "Create a Google Form or spreadsheet for native-speaker scoring.", "Next"],
            ["3", "Collect at least 5 pilot annotations before scaling to more annotators.", "Next"],
            ["4", "Implement model evaluation once the first scored CSV is available.", "Next"],
        ],
        [0.65, 4.7, 1.0],
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
